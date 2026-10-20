"""Transaction ingestion, trade parsing, union, and deduplication (Task 2)."""

import csv
import json
import re
import sys
from collections.abc import Sequence
from datetime import date
from pathlib import Path
from typing import Any

from commsec_pairing.matcher import get_financial_year
from commsec_pairing.models import ParsedTrade, TradeAction, Transaction

TRADE_PATTERN = re.compile(
    r"^\s*(?P<action>[BS])\s+(?P<units>\d+)\s+(?P<security>[A-Z0-9_.-]+)\s*@\s*(?P<price>[\d\.]+)",
    re.IGNORECASE,
)



def parse_date_str(date_str: str) -> date:
    """Parse date strings supporting DD/MM/YYYY and YYYY-MM-DD formats."""
    cleaned = date_str.strip()
    if "/" in cleaned:
        day_str, month_str, year_str = cleaned.split("/")
        return date(int(year_str), int(month_str), int(day_str))
    elif "-" in cleaned:
        return date.fromisoformat(cleaned)
    raise ValueError(f"Unsupported date format: '{date_str}'")


def parse_trade_details(details_str: str) -> ParsedTrade | None:
    """Extract trade action, quantity, security ticker, and price from details string."""
    if not details_str:
        return None
    match = TRADE_PATTERN.search(details_str.strip())
    if not match:
        return None

    action_char = match.group("action").upper()
    action = TradeAction.BUY if action_char == "B" else TradeAction.SELL
    units = int(match.group("units"))
    security = match.group("security").upper()
    unit_price = float(match.group("price"))

    return ParsedTrade(
        action=action,
        units=units,
        security=security,
        unit_price=unit_price,
    )


def _clean_float(val: str | None) -> float | None:
    if not val:
        return None
    cleaned = val.strip().replace("$", "").replace(",", "")
    if not cleaned:
        return None
    return float(cleaned)


def load_transactions_csv(file_path: str | Path) -> list[Transaction]:
    """Load and parse transactions from a single CommSec CSV file."""
    path = Path(file_path)
    if not path.is_file():
        raise FileNotFoundError(f"Transaction CSV file not found: {path}")

    transactions: list[Transaction] = []
    with open(path, mode="r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for row in reader:
            normalized_row = {k.strip(): v for k, v in row.items() if k is not None}

            raw_date = normalized_row.get("Date", "")
            if not raw_date:
                continue

            parsed_d = parse_date_str(raw_date)
            reference = normalized_row.get("Reference", "").strip()
            details = normalized_row.get("Details", "").strip()
            debit = _clean_float(normalized_row.get("Debit($)") or normalized_row.get("Debit"))
            credit = _clean_float(normalized_row.get("Credit($)") or normalized_row.get("Credit"))
            balance = _clean_float(
                normalized_row.get("Balance($)") or normalized_row.get("Balance")
            )

            raw_is_short = normalized_row.get("IsShort") or normalized_row.get("is_short")
            is_short = (
                str(raw_is_short).strip().lower() in ("true", "1", "yes")
                if raw_is_short
                else False
            )

            parsed_trade = parse_trade_details(details) if reference.startswith("C") else None
            if parsed_trade and is_short:
                parsed_trade.is_short = True

            transactions.append(
                Transaction(
                    date=parsed_d,
                    reference=reference,
                    details=details,
                    debit=debit,
                    credit=credit,
                    balance=balance,
                    parsed_trade=parsed_trade,
                    is_short=is_short,
                )
            )

    return transactions


def discover_transaction_csvs(entity_dir: str | Path) -> list[Path]:
    """Find all transaction CSV files in an entity directory."""
    path = Path(entity_dir)
    if not path.is_dir():
        raise NotADirectoryError(f"Entity directory does not exist: {path}")

    csvs = sorted(path.glob("Transactions_*.csv"))
    if not csvs:
        csvs = sorted(path.glob("[Tt]ransactions*.csv"))
    return csvs


def get_deduplication_fingerprint(tx: Transaction) -> tuple:
    """
    Generate deduplication fingerprint for a transaction.
    For trade contracts, the fingerprint strictly consists of:
    (contract reference, date, security ticker, number of units, action B/S).
    """
    if tx.is_contract and tx.parsed_trade:
        return (
            tx.reference,
            tx.date,
            tx.parsed_trade.security,
            tx.parsed_trade.units,
            tx.parsed_trade.action,
        )
    return (
        tx.reference,
        tx.date,
        tx.details,
        tx.debit,
        tx.credit,
    )


def union_and_deduplicate(csv_paths: Sequence[str | Path]) -> list[Transaction]:
    """Union multiple transaction CSV files and deduplicate records by fingerprint."""
    seen = set()
    unified: list[Transaction] = []

    for path in csv_paths:
        txs = load_transactions_csv(path)
        for tx in txs:
            key = get_deduplication_fingerprint(tx)
            if key not in seen:
                seen.add(key)
                unified.append(tx)

    unified.sort(key=lambda x: (x.date, x.reference))
    return unified



def save_transactions_to_csv(
    transactions: Sequence[Transaction], output_path: str | Path
) -> None:
    """Save normalized transactions to CSV format."""
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)

    has_short = any(
        getattr(tx, "is_short", False)
        or (tx.parsed_trade and getattr(tx.parsed_trade, "is_short", False))
        for tx in transactions
    )

    with open(out, mode="w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        if has_short:
            writer.writerow(["Date", "Reference", "Details", "Debit($)", "Credit($)", "Balance($)", "IsShort"])
            for tx in transactions:
                writer.writerow([
                    tx.date.strftime("%d/%m/%Y"),
                    tx.reference,
                    tx.details,
                    f"{tx.debit:.2f}" if tx.debit is not None else "",
                    f"{tx.credit:.2f}" if tx.credit is not None else "",
                    f"{tx.balance:.2f}" if tx.balance is not None else "",
                    str(tx.is_short),
                ])
        else:
            writer.writerow(["Date", "Reference", "Details", "Debit($)", "Credit($)", "Balance($)"])
            for tx in transactions:
                writer.writerow([
                    tx.date.strftime("%d/%m/%Y"),
                    tx.reference,
                    tx.details,
                    f"{tx.debit:.2f}" if tx.debit is not None else "",
                    f"{tx.credit:.2f}" if tx.credit is not None else "",
                    f"{tx.balance:.2f}" if tx.balance is not None else "",
                ])


def load_options_resolutions(resolution_file: str | Path) -> dict[str, dict]:
    """Load cross-FY options resolutions from JSON file."""
    path = Path(resolution_file)
    if not path.is_file():
        return {}
    with open(path, mode="r", encoding="utf-8") as f:
        try:
            data = json.load(f)
            return data if isinstance(data, dict) else {}
        except json.JSONDecodeError:
            return {}


def save_options_resolutions(
    resolutions: dict[str, dict], resolution_file: str | Path
) -> None:
    """Save cross-FY options resolutions to JSON file."""
    path = Path(resolution_file)
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, mode="w", encoding="utf-8") as f:
        json.dump(resolutions, f, indent=2)


def prompt_option_resolution(
    code: str,
    s_date: date,
    exp_date: date,
    units: int,
    open_price: float,
) -> dict[str, Any]:
    """
    Prompt user interactively for the outcome of an option position open across the FY boundary.
    Outcomes according to ATO CGT rules:
    1) Expired worthless (CGT event D2: cost base $0.00, proceeds = open price)
    2) Bought back / Closed (cost base = buy back price, proceeds = open price)
    3) Exercised (CGT event D2 disregarded; rolled into underlying stock disposal/acquisition)
    """
    from rich.console import Console
    from rich.prompt import Prompt

    cons = Console(file=sys.stderr)
    cons.print(
        f"\n[bold yellow]Cross-FY Option Position Detected:[/] [bold cyan]{code}[/]\n"
        f"  • Open Leg: Written on {s_date.strftime('%d/%m/%Y')} @ ${open_price:.4f} ({units} units / {units // 100} contract(s))\n"
        f"  • Expiry Date: {exp_date.strftime('%d/%m/%Y')} (Crosses FY boundary)"
    )
    cons.print("[dim]Under ATO rules (CGT Event D2), choose how this position resolved:[/]")
    cons.print("  [1] Expired worthless (Cost base $0.00, realized capital gain in grant FY)")
    cons.print("  [2] Closed / Bought back (Enter closing buy-back price and date)")
    cons.print("  [3] Exercised (Disregarded in grant FY, rolled into underlying shares)")

    choice = Prompt.ask(
        "Select outcome",
        choices=["1", "2", "3", "expired", "closed", "exercised"],
        default="1",
        console=cons,
    )

    if choice in ("1", "expired"):
        return {
            "status": "expired",
            "close_price": 0.0,
            "close_date": exp_date.isoformat(),
            "notes": "Expired worthless",
        }
    elif choice in ("2", "closed"):
        price_str = Prompt.ask("Enter closing buy-back price per share (USD)", console=cons)
        close_price = float(price_str)
        close_date_str = Prompt.ask(
            f"Enter closing date [DD/MM/YYYY or YYYY-MM-DD] (default {exp_date.strftime('%d/%m/%Y')})",
            default=exp_date.strftime("%d/%m/%Y"),
            console=cons,
        )
        parsed_close_d = parse_date_str(close_date_str)
        return {
            "status": "closed",
            "close_price": close_price,
            "close_date": parsed_close_d.isoformat(),
            "notes": f"Bought back at {close_price}",
        }
    else:
        return {
            "status": "exercised",
            "close_price": 0.0,
            "close_date": exp_date.isoformat(),
            "notes": "Exercised into underlying stock",
        }


OCC_OPTION_PATTERN = re.compile(r"^([A-Z]+)(\d{2})(\d{2})(\d{2})([CP])(\d+)$")


def discover_moomoo_excel_files(entity_dir: str | Path) -> list[Path]:
    """Find all MooMoo Excel (.xlsx) files in an entity directory."""
    path = Path(entity_dir)
    if not path.is_dir():
        raise NotADirectoryError(f"Entity directory does not exist: {path}")

    return sorted(path.glob("*.xlsx"))


def convert_moomoo_excel_to_transactions(
    excel_path: str | Path,
    resolution_file: str | Path | None = None,
    resolutions: dict[str, dict] | None = None,
    interactive: bool = True,
) -> list[Transaction]:
    """
    Parse a MooMoo Financial Year End Statement (.xlsx), extracting Stock and Option
    trades from the 'Transaction Overview' sheet, strictly ignoring dividends, interest, and cash.

    Follows Australian Taxation Office (ATO) rules for options (CGT Event D2 / C2):
    - Expired written options (short sell without buy-back): cost base = 0.00, proceeds = short sell price.
      Under strict CGT Event D2, disposal is recognized on the grant trade date.
    - Explicitly closed written options (buy-back): cost base = buy-back price, proceeds = short sell price.
      Recognized on the closing buy-back date.
    - Expired long options (bought without sell): cost base = purchase price, proceeds = 0.00.
      Disposal recognized on the expiry date (CGT Event C2).
    - Primary key for options: OCC 3-tuple identifier (Underlying, Expiration, Strike & Type).
    - 1 US option contract = 100 shares.
    """
    import openpyxl

    path = Path(excel_path)
    if not path.is_file():
        raise FileNotFoundError(f"MooMoo Excel file not found: {path}")

    res_path = (
        Path(resolution_file)
        if resolution_file
        else Path(excel_path).parent / "options_resolution.json"
    )
    loaded_resolutions: dict[str, dict] = (
        dict(resolutions)
        if resolutions is not None
        else load_options_resolutions(res_path)
    )

    wb = openpyxl.load_workbook(path, data_only=True)
    if "Transaction Overview" not in wb.sheetnames:
        raise ValueError(
            f"Expected 'Transaction Overview' sheet in MooMoo Excel file: {path.name}"
        )

    ws = wb["Transaction Overview"]
    headers = [ws.cell(row=1, column=j).value for j in range(1, ws.max_column + 1)]
    header_indices = {
        str(h).strip(): idx for idx, h in enumerate(headers) if h is not None
    }

    required_cols = [
        "Security Type",
        "Direction",
        "Security Code",
        "Quantity",
        "Avg Price",
        "Transaction Date",
    ]
    for col in required_cols:
        if col not in header_indices:
            raise ValueError(
                f"Missing required column '{col}' in MooMoo sheet 'Transaction Overview'"
            )

    raw_rows = []
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(row=r, column=j).value for j in range(1, ws.max_column + 1)]
        if any(v is not None for v in vals):
            row_dict = {
                k: vals[idx] if idx < len(vals) else None
                for k, idx in header_indices.items()
            }
            raw_rows.append(row_dict)

    transactions: list[Transaction] = []
    ref_idx = 1000

    # 1. Stocks
    stock_rows = [
        r for r in raw_rows if str(r.get("Security Type", "")).strip().lower() == "stock"
    ]
    for r in stock_rows:
        raw_date = r.get("Transaction Date")
        if not raw_date:
            continue
        dt = parse_date_str(str(raw_date)) if isinstance(raw_date, str) else raw_date
        if hasattr(dt, "date"):
            dt = dt.date()

        sec = str(r.get("Security Code", "")).strip().upper()
        qty = int(float(r.get("Quantity") or 0))
        price = float(r.get("Avg Price") or 0.0)
        direction = str(r.get("Direction", "")).strip()

        action = TradeAction.BUY if direction in ("Buy", "Buy Back") else TradeAction.SELL
        details = f"{action.value} {qty} {sec} @ {price:.4f}"
        raw_amount = r.get("Transaction Amount")
        amt = abs(float(raw_amount)) if raw_amount is not None else round(qty * price, 2)
        debit = amt if action == TradeAction.BUY else None
        credit = amt if action == TradeAction.SELL else None

        ref = f"C{ref_idx}"
        ref_idx += 1

        transactions.append(
            Transaction(
                date=dt,
                reference=ref,
                details=details,
                debit=debit,
                credit=credit,
                parsed_trade=ParsedTrade(
                    action=action,
                    units=qty,
                    security=sec,
                    unit_price=price,
                    is_short=False,
                ),
                is_short=False,
            )
        )

    # 2. Options
    option_rows = [
        r for r in raw_rows if str(r.get("Security Type", "")).strip().lower() == "option"
    ]
    from collections import defaultdict

    options_by_code: dict[str, list[dict]] = defaultdict(list)
    for r in option_rows:
        code = str(r.get("Security Code", "")).strip().upper()
        if code:
            options_by_code[code].append(r)

    for code, txs in options_by_code.items():
        m = OCC_OPTION_PATTERN.match(code)
        if not m:
            continue
        _, yy, mm, dd, _, _ = m.groups()
        exp_date = date(2000 + int(yy), int(mm), int(dd))

        sells: list[dict] = []
        buy_backs: list[dict] = []
        buys: list[dict] = []

        for t in txs:
            d = str(t.get("Direction", "")).strip()
            qty = int(float(t.get("Quantity") or 0)) * 100  # 1 US option contract = 100 shares
            p = float(t.get("Avg Price") or 0.0)
            raw_d = t.get("Transaction Date")
            t_date = parse_date_str(str(raw_d)) if isinstance(raw_d, str) else raw_d
            if hasattr(t_date, "date"):
                t_date = t_date.date()

            item = {"date": t_date, "price": p, "units": qty}
            if d == "Sell Short":
                sells.append(item)
            elif d == "Buy Back":
                buy_backs.append(item)
            elif d == "Buy":
                buys.append(item)

        # Explicit close (buy back within statement)
        for bb in buy_backs:
            while bb["units"] > 0 and sells:
                s = sells[0]
                matched = min(s["units"], bb["units"])
                ref_b = f"C{ref_idx}"
                ref_idx += 1
                transactions.append(
                    Transaction(
                        date=s["date"],
                        reference=ref_b,
                        details=f"B {matched} {code} @ {bb['price']:.4f}",
                        debit=round(matched * bb["price"], 2),
                        credit=None,
                        parsed_trade=ParsedTrade(
                            action=TradeAction.BUY,
                            units=matched,
                            security=code,
                            unit_price=bb["price"],
                            is_short=True,
                        ),
                        is_short=True,
                    )
                )
                ref_s = f"C{ref_idx}"
                ref_idx += 1
                transactions.append(
                    Transaction(
                        date=bb["date"],
                        reference=ref_s,
                        details=f"S {matched} {code} @ {s['price']:.4f}",
                        debit=None,
                        credit=round(matched * s["price"], 2),
                        parsed_trade=ParsedTrade(
                            action=TradeAction.SELL,
                            units=matched,
                            security=code,
                            unit_price=s["price"],
                            is_short=True,
                        ),
                        is_short=True,
                    )
                )
                s["units"] -= matched
                bb["units"] -= matched
                if s["units"] == 0:
                    sells.pop(0)

        # Unclosed written options (either boundary or expired in same FY)
        for s in sells:
            if s["units"] <= 0:
                continue
            is_boundary = get_financial_year(exp_date) > get_financial_year(s["date"])
            if is_boundary:
                # Boundary option
                if code in loaded_resolutions:
                    res = loaded_resolutions[code]
                elif interactive and sys.stdin.isatty():
                    res = prompt_option_resolution(
                        code=code,
                        s_date=s["date"],
                        exp_date=exp_date,
                        units=s["units"],
                        open_price=s["price"],
                    )
                    loaded_resolutions[code] = res
                    save_options_resolutions(loaded_resolutions, res_path)
                else:
                    raise ValueError(
                        f"Unresolved cross-FY option position: {code} (written {s['date']}, exp {exp_date}). "
                        f"Please provide resolution via interactive prompt or in {res_path.name}"
                    )

                status = str(res.get("status", "expired")).lower()
                if status == "exercised":
                    # Disregarded under CGT Event D2 (s 104-40(5) ITAA 1997)
                    continue

                close_price = float(res.get("close_price", 0.0))
                ref_b = f"C{ref_idx}"
                ref_idx += 1
                transactions.append(
                    Transaction(
                        date=s["date"],
                        reference=ref_b,
                        details=f"B {s['units']} {code} @ {close_price:.4f}",
                        debit=round(s["units"] * close_price, 2) if close_price > 0 else 0.0,
                        credit=None,
                        parsed_trade=ParsedTrade(
                            action=TradeAction.BUY,
                            units=s["units"],
                            security=code,
                            unit_price=close_price,
                            is_short=True,
                        ),
                        is_short=True,
                    )
                )
                ref_s = f"C{ref_idx}"
                ref_idx += 1
                transactions.append(
                    Transaction(
                        date=s["date"],
                        reference=ref_s,
                        details=f"S {s['units']} {code} @ {s['price']:.4f}",
                        debit=None,
                        credit=round(s["units"] * s["price"], 2),
                        parsed_trade=ParsedTrade(
                            action=TradeAction.SELL,
                            units=s["units"],
                            security=code,
                            unit_price=s["price"],
                            is_short=True,
                        ),
                        is_short=True,
                    )
                )
            else:
                # Expired written options within the same FY (strict ATO D2 model)
                ref_b = f"C{ref_idx}"
                ref_idx += 1
                transactions.append(
                    Transaction(
                        date=s["date"],
                        reference=ref_b,
                        details=f"B {s['units']} {code} @ 0.0000",
                        debit=0.0,
                        credit=None,
                        parsed_trade=ParsedTrade(
                            action=TradeAction.BUY,
                            units=s["units"],
                            security=code,
                            unit_price=0.0,
                            is_short=True,
                        ),
                        is_short=True,
                    )
                )
                ref_s = f"C{ref_idx}"
                ref_idx += 1
                transactions.append(
                    Transaction(
                        date=s["date"],
                        reference=ref_s,
                        details=f"S {s['units']} {code} @ {s['price']:.4f}",
                        debit=None,
                        credit=round(s["units"] * s["price"], 2),
                        parsed_trade=ParsedTrade(
                            action=TradeAction.SELL,
                            units=s["units"],
                            security=code,
                            unit_price=s["price"],
                            is_short=True,
                        ),
                        is_short=True,
                    )
                )

        # Long options
        for b in buys:
            if b["units"] > 0:
                is_boundary = get_financial_year(exp_date) > get_financial_year(b["date"])
                if is_boundary:
                    # Long option is open across FY boundary; carried forward as buy inventory
                    ref_b = f"C{ref_idx}"
                    ref_idx += 1
                    transactions.append(
                        Transaction(
                            date=b["date"],
                            reference=ref_b,
                            details=f"B {b['units']} {code} @ {b['price']:.4f}",
                            debit=round(b["units"] * b["price"], 2),
                            credit=None,
                            parsed_trade=ParsedTrade(
                                action=TradeAction.BUY,
                                units=b["units"],
                                security=code,
                                unit_price=b["price"],
                                is_short=False,
                            ),
                            is_short=False,
                        )
                    )
                else:
                    # Expired in same FY (CGT event C2 happens on expiry date)
                    ref_b = f"C{ref_idx}"
                    ref_idx += 1
                    transactions.append(
                        Transaction(
                            date=b["date"],
                            reference=ref_b,
                            details=f"B {b['units']} {code} @ {b['price']:.4f}",
                            debit=round(b["units"] * b["price"], 2),
                            credit=None,
                            parsed_trade=ParsedTrade(
                                action=TradeAction.BUY,
                                units=b["units"],
                                security=code,
                                unit_price=b["price"],
                                is_short=False,
                            ),
                            is_short=False,
                        )
                    )
                    ref_s = f"C{ref_idx}"
                    ref_idx += 1
                    transactions.append(
                        Transaction(
                            date=exp_date,
                            reference=ref_s,
                            details=f"S {b['units']} {code} @ 0.0000",
                            debit=None,
                            credit=0.0,
                            parsed_trade=ParsedTrade(
                                action=TradeAction.SELL,
                                units=b["units"],
                                security=code,
                                unit_price=0.0,
                                is_short=False,
                            ),
                            is_short=False,
                        )
                    )

    transactions.sort(key=lambda x: (x.date, x.reference))
    return transactions


def export_moomoo_to_csv(
    excel_path: str | Path,
    output_csv_path: str | Path,
    resolution_file: str | Path | None = None,
    resolutions: dict[str, dict] | None = None,
    interactive: bool = True,
) -> Path:
    """Parse a MooMoo Excel statement and save normalized transactions to a standard CSV."""
    txs = convert_moomoo_excel_to_transactions(
        excel_path,
        resolution_file=resolution_file,
        resolutions=resolutions,
        interactive=interactive,
    )
    out_path = Path(output_csv_path)
    save_transactions_to_csv(txs, out_path)
    return out_path


