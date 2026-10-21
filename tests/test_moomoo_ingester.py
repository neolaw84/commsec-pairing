"""Unit tests for MooMoo statement ingestion and option trading calculations."""

from pathlib import Path

import openpyxl
import pytest

from commsec_pairing.exporter import export_reports
from commsec_pairing.ingester import (
    convert_moomoo_excel_to_transactions,
    export_moomoo_to_csv,
    load_options_resolutions,
    load_transactions_csv,
    parse_trade_details,
    save_options_resolutions,
)
from commsec_pairing.matcher import CGTMatcher
from commsec_pairing.models import PersistentMappingStore, TradeAction
from commsec_pairing.storage import MappingStorageManager


def _create_synthetic_moomoo_workbook(file_path: Path) -> Path:
    """Create a synthetic MooMoo-format Excel workbook with stock and option trades."""
    wb = openpyxl.Workbook()

    # Sheet 1: Account Overview (dummy metadata)
    ws_acc = wb.active
    ws_acc.title = "Account Overview"
    ws_acc.append(["Name", "Moomoo ID", "Account ID", "HIN", "Address"])
    ws_acc.append(["Test User", 12345678, "99999999", "X0000000000", "Sydney NSW 2000"])

    # Sheet 2: Summary (dummy metadata)
    ws_sum = wb.create_sheet(title="Summary")
    ws_sum.append(["Dividends", "Currency", "Value"])
    ws_sum.append(["Dividend", "USD", 100.0])

    # Sheet 3: Transaction Overview (In-scope trades)
    ws_tx = wb.create_sheet(title="Transaction Overview")
    ws_tx.append([
        "Market",
        "Transaction Date",
        "Settlement Date",
        "Security Type",
        "Direction",
        "Security Code",
        "Security Name",
        "Quantity",
        "Currency",
        "Avg Price",
        "Transaction Amount",
        "Transaction Fee(Inc.GST)",
        "GST",
        "Net Amount",
        "Comment",
    ])

    # Row 1: Stock Buy
    ws_tx.append([
        "US", "10/08/2025", "11/08/2025", "Stock", "Buy", "XYZ", "XYZ Corp", 10, "USD", 50.0, -500.0, 1.0, 0, -501.0, ""
    ])
    # Row 2: Stock Sell
    ws_tx.append([
        "US", "15/10/2025", "16/10/2025", "Stock", "Sell", "XYZ", "XYZ Corp", 10, "USD", 65.0, 650.0, 1.0, 0, 649.0, ""
    ])

    # Row 3: Option Written and Expired (Short Sale without buy-back)
    # OCC: ABC 2025-08-15 Put 10.00
    ws_tx.append([
        "US", "15/07/2025", "16/07/2025", "Option", "Sell Short", "ABC250815P10000", "ABC 250815 10.00P", 1, "USD", 0.50, 50.0, 0.50, 0, 49.50, ""
    ])

    # Row 4: Option Written and Explicitly Closed (Sell Short + Buy Back)
    # OCC: DEF 2025-09-19 Call 25.00
    ws_tx.append([
        "US", "01/08/2025", "02/08/2025", "Option", "Sell Short", "DEF250919C25000", "DEF 250919 25.00C", 1, "USD", 1.20, 120.0, 0.50, 0, 119.50, ""
    ])
    ws_tx.append([
        "US", "20/08/2025", "21/08/2025", "Option", "Buy Back", "DEF250919C25000", "DEF 250919 25.00C", 1, "USD", 0.30, -30.0, 0.50, 0, -30.50, ""
    ])

    # Row 5: Long Option Bought and Expired Worthless
    # OCC: GHI 2025-08-01 Put 5.00
    ws_tx.append([
        "US", "12/07/2025", "13/07/2025", "Option", "Buy", "GHI250801P5000", "GHI 250801 5.00P", 1, "USD", 0.20, -20.0, 0.50, 0, -20.50, ""
    ])

    wb.save(file_path)
    return file_path


def test_parse_trade_details_with_option_code():
    """Verify that TRADE_PATTERN parses OCC option codes properly."""
    parsed = parse_trade_details("B 100 ABC250815P10000 @ 0.5000")
    assert parsed is not None
    assert parsed.action == TradeAction.BUY
    assert parsed.units == 100
    assert parsed.security == "ABC250815P10000"
    assert parsed.unit_price == 0.50

    parsed_sell = parse_trade_details("S 100 DEF250919C25000 @ 1.2000")
    assert parsed_sell is not None
    assert parsed_sell.action == TradeAction.SELL
    assert parsed_sell.units == 100
    assert parsed_sell.security == "DEF250919C25000"
    assert parsed_sell.unit_price == 1.20


def test_convert_moomoo_excel_to_transactions(tmp_path: Path):
    """Test full conversion of synthetic MooMoo statement to transactions."""
    excel_file = tmp_path / "synthetic_moomoo.xlsx"
    _create_synthetic_moomoo_workbook(excel_file)

    txs = convert_moomoo_excel_to_transactions(excel_file)
    assert len(txs) > 0

    # Stock checks
    xyz_txs = [t for t in txs if t.parsed_trade and t.parsed_trade.security == "XYZ"]
    assert len(xyz_txs) == 2
    assert xyz_txs[0].parsed_trade.action == TradeAction.BUY
    assert xyz_txs[0].parsed_trade.units == 10
    assert xyz_txs[0].parsed_trade.unit_price == 50.0
    assert xyz_txs[1].parsed_trade.action == TradeAction.SELL
    assert xyz_txs[1].parsed_trade.units == 10
    assert xyz_txs[1].parsed_trade.unit_price == 65.0

    # Option: ABC250815P10000 (Expired short)
    # 1 contract = 100 shares. Cost base = 0.0, proceeds = 0.50
    abc_txs = [t for t in txs if t.parsed_trade and t.parsed_trade.security == "ABC250815P10000"]
    assert len(abc_txs) == 2
    b_leg = next(t for t in abc_txs if t.parsed_trade.action == TradeAction.BUY)
    s_leg = next(t for t in abc_txs if t.parsed_trade.action == TradeAction.SELL)
    assert b_leg.parsed_trade.units == 100
    assert b_leg.parsed_trade.unit_price == 0.0
    assert s_leg.parsed_trade.units == 100
    assert s_leg.parsed_trade.unit_price == 0.50

    # Option: DEF250919C25000 (Explicitly closed)
    # Sold @ 1.20, bought back @ 0.30
    def_txs = [t for t in txs if t.parsed_trade and t.parsed_trade.security == "DEF250919C25000"]
    assert len(def_txs) == 2
    def_b = next(t for t in def_txs if t.parsed_trade.action == TradeAction.BUY)
    def_s = next(t for t in def_txs if t.parsed_trade.action == TradeAction.SELL)
    assert def_b.parsed_trade.units == 100
    assert def_b.parsed_trade.unit_price == 0.30
    assert def_s.parsed_trade.units == 100
    assert def_s.parsed_trade.unit_price == 1.20

    # Option: GHI250801P5000 (Expired long)
    # Bought @ 0.20, expired @ 0.0
    ghi_txs = [t for t in txs if t.parsed_trade and t.parsed_trade.security == "GHI250801P5000"]
    assert len(ghi_txs) == 2
    ghi_b = next(t for t in ghi_txs if t.parsed_trade.action == TradeAction.BUY)
    ghi_s = next(t for t in ghi_txs if t.parsed_trade.action == TradeAction.SELL)
    assert ghi_b.parsed_trade.units == 100
    assert ghi_b.parsed_trade.unit_price == 0.20
    assert ghi_s.parsed_trade.units == 100
    assert ghi_s.parsed_trade.unit_price == 0.0


def test_export_moomoo_to_csv_and_reload(tmp_path: Path):
    """Test exporting to CSV and re-loading via standard transactions parser."""
    excel_file = tmp_path / "synthetic_moomoo.xlsx"
    csv_file = tmp_path / "Transactions_moomoo.csv"
    _create_synthetic_moomoo_workbook(excel_file)

    export_moomoo_to_csv(excel_file, csv_file)
    assert csv_file.is_file()

    reloaded = load_transactions_csv(csv_file)
    assert len(reloaded) > 0
    for tx in reloaded:
        assert tx.is_contract
        assert tx.parsed_trade is not None


def test_moomoo_end_to_end_matching(tmp_path: Path):
    """Test full matching workflow on synthetic MooMoo statement data."""
    excel_file = tmp_path / "synthetic_moomoo.xlsx"
    csv_file = tmp_path / "Transactions_moomoo.csv"
    _create_synthetic_moomoo_workbook(excel_file)

    export_moomoo_to_csv(excel_file, csv_file)
    txs = load_transactions_csv(csv_file)

    store = PersistentMappingStore(entity="test-entity")
    mgr = MappingStorageManager(tmp_path / "mapping.json", "test-entity")
    store = mgr.sync_inventory_from_transactions(store, txs, max_fy=2026)

    matcher = CGTMatcher(store)
    matches = matcher.optimize_financial_year(txs, target_fy=2026)

    # 4 disposals in FY2026:
    # 1. Stock XYZ: 10 * (65 - 50) = +150.00
    # 2. Option ABC (expired written): 100 * (0.50 - 0.0) = +50.00
    # 3. Option DEF (closed short): 100 * (1.20 - 0.30) = +90.00
    # 4. Option GHI (expired long): 100 * (0.0 - 0.20) = -20.00
    # Total net taxable gain: 150 + 50 + 90 - 20 = 270.00
    assert len(matches) == 4
    total_taxable = sum(m.taxable_gain_loss_aud for m in matches)
    assert round(total_taxable, 2) == 270.00


def test_cross_fy_boundary_options_resolution(tmp_path: Path):
    """Test resolution of options crossing the Australian FY boundary (expired, closed, exercised)."""
    excel_file = tmp_path / "boundary_moomoo.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Transaction Overview"
    ws.append([
        "Market", "Transaction Date", "Settlement Date", "Security Type",
        "Direction", "Security Code", "Security Name", "Quantity",
        "Currency", "Avg Price", "Transaction Amount", "Transaction Fee(Inc.GST)",
        "GST", "Net Amount", "Comment"
    ])
    # Option 1: Written late June 2026, expires 24 July 2026 (FY2027) -> Expired worthless
    ws.append([
        "US", "27/06/2026", "28/06/2026", "Option", "Sell Short", "TEST260724C16000", "TEST Call",
        1, "USD", 0.24, 24.0, 0.50, 0, 23.50, ""
    ])
    # Option 2: Written late June 2026, expires 17 July 2026 (FY2027) -> Closed at 0.29 (loss)
    ws.append([
        "US", "29/06/2026", "30/06/2026", "Option", "Sell Short", "TEST260717C14000", "TEST Call",
        1, "USD", 0.11, 11.0, 0.50, 0, 10.50, ""
    ])
    # Option 3: Written late June 2026, expires 31 July 2026 (FY2027) -> Exercised (rolled over)
    ws.append([
        "US", "30/06/2026", "01/07/2026", "Option", "Sell Short", "TEST260731C20000", "TEST Call",
        1, "USD", 0.50, 50.0, 0.50, 0, 49.50, ""
    ])
    wb.save(excel_file)

    # 1. Unresolved boundary options without interactive input should raise ValueError
    with pytest.raises(ValueError, match="Unresolved cross-FY option position"):
        convert_moomoo_excel_to_transactions(excel_file, interactive=False)

    # 2. Persist resolution file
    res_file = tmp_path / "options_resolution.json"
    resolutions = {
        "TEST260724C16000": {
            "status": "expired",
            "close_price": 0.0,
            "close_date": "2026-07-24",
            "notes": "Expired worthless"
        },
        "TEST260717C14000": {
            "status": "closed",
            "close_price": 0.29,
            "close_date": "2026-07-17",
            "notes": "Closed at 0.29"
        },
        "TEST260731C20000": {
            "status": "exercised",
            "close_price": 0.0,
            "close_date": "2026-07-31",
            "notes": "Exercised into stock"
        }
    }
    save_options_resolutions(resolutions, res_file)
    loaded = load_options_resolutions(res_file)
    assert loaded == resolutions

    # 3. Convert with resolutions
    txs = convert_moomoo_excel_to_transactions(excel_file, resolution_file=res_file, interactive=False)
    # TEST260724C16000 (expired): 2 legs (Buy @ 0.0, Sell @ 0.24)
    # TEST260717C14000 (closed): 2 legs (Buy @ 0.29, Sell @ 0.11)
    # TEST260731C20000 (exercised): 0 legs (disregarded in grant FY under s 104-40(5))
    assert len(txs) == 4
    for t in txs:
        assert t.is_short is True
        assert t.parsed_trade.is_short is True

    # 4. Matching & Reports
    store = PersistentMappingStore(entity="test-boundary")
    mgr = MappingStorageManager(tmp_path / "mapping.json", "test-boundary")
    store = mgr.sync_inventory_from_transactions(store, txs, max_fy=2026)
    matcher = CGTMatcher(store)
    matches = matcher.optimize_financial_year(txs, target_fy=2026)
    assert len(matches) == 2

    m_exp = next(m for m in matches if m.security == "TEST260724C16000")
    m_cls = next(m for m in matches if m.security == "TEST260717C14000")

    assert m_exp.is_short is True
    assert round(m_exp.taxable_gain_loss_aud, 2) == 24.00  # 100 * (0.24 - 0.0)

    assert m_cls.is_short is True
    assert round(m_cls.taxable_gain_loss_aud, 2) == -18.00  # 100 * (0.11 - 0.29)

    for m in matches:
        mgr.record_match(store, m)

    # 5. Export Reports and verify formatting
    reports = export_reports(store, output_dir=tmp_path / "reports")
    parcels_content = reports["parcels_csv"].read_text(encoding="utf-8")
    assert "PositionType" in parcels_content
    assert "Short" in parcels_content

    audit_md = reports["report_md"].read_text(encoding="utf-8")
    assert "Sell Short:" in audit_md
    assert "Expired ($0.00)" in audit_md
    assert "Buy Back:" in audit_md

